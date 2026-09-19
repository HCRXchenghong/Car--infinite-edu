"""Build a Chinese A4 review handbook and editable DOCX from controlled Markdown.

Dependencies: reportlab, python-docx, pypdf. Outputs are generated artifacts.
Set --font to an embeddable Chinese TrueType font on another workstation.
"""
from pathlib import Path
import argparse
import hashlib
import html
import json
import re

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import BaseDocTemplate, Frame, PageTemplate, Paragraph, Spacer, PageBreak, CondPageBreak, LongTable, TableStyle
from reportlab.platypus.tableofcontents import TableOfContents
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / '打印版'
OUT.mkdir(exist_ok=True)
parser = argparse.ArgumentParser()
parser.add_argument('--font', default='/System/Library/Fonts/Supplemental/Arial Unicode.ttf')
args = parser.parse_args()
assert Path(args.font).is_file(), 'Provide --font with a Chinese TrueType font path.'
pdfmetrics.registerFont(TTFont('Chinese', args.font))
pdfmetrics.registerFontFamily('Chinese', normal='Chinese', bold='Chinese', italic='Chinese', boldItalic='Chinese')

FILES=sorted(ROOT.glob('[0-9][0-9]_*.md'))
assert len(FILES)==12, [p.name for p in FILES]
blocks=[]
for path in FILES:
    lines=path.read_text(encoding='utf-8').splitlines()
    i=0
    while i<len(lines):
        line=lines[i].strip()
        if not line: i+=1; continue
        if line.startswith('```'):
            code=[]; i+=1
            while i<len(lines) and not lines[i].strip().startswith('```'):
                code.append(lines[i]); i+=1
            blocks.append(('code','\n'.join(code),path.name)); i+=1; continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].strip().startswith('|'):
                cells=[c.strip() for c in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-{3,}:?',c) for c in cells): rows.append(cells)
                i+=1
            assert rows and all(len(r)==len(rows[0]) for r in rows),path.name
            blocks.append(('table',rows,path.name)); continue
        m=re.match(r'^(#{1,3})\s+(.+)',line)
        if m: blocks.append(('h'+str(len(m[1])),m[2],path.name))
        else: blocks.append(('p',line,path.name))
        i+=1

NAVY=colors.HexColor('#143342'); TEAL=colors.HexColor('#186C73'); GRAY=colors.HexColor('#52636C')
styles={
 'body':ParagraphStyle('body',fontName='Chinese',fontSize=9.5,leading=15,spaceAfter=6,wordWrap='CJK',textColor=NAVY),
 'h1':ParagraphStyle('h1',fontName='Chinese',fontSize=20,leading=28,spaceAfter=15,textColor=NAVY,keepWithNext=True,wordWrap='CJK'),
 'h2':ParagraphStyle('h2',fontName='Chinese',fontSize=12.5,leading=19,spaceBefore=13,spaceAfter=7,textColor=TEAL,keepWithNext=True,wordWrap='CJK'),
 'h3':ParagraphStyle('h3',fontName='Chinese',fontSize=10.5,leading=16,spaceBefore=9,spaceAfter=6,textColor=TEAL,keepWithNext=True,wordWrap='CJK'),
 'cell':ParagraphStyle('cell',fontName='Chinese',fontSize=8,leading=11.5,wordWrap='CJK',textColor=NAVY),
 'thead':ParagraphStyle('thead',fontName='Chinese',fontSize=8.3,leading=12,wordWrap='CJK',textColor=colors.white),
 'code':ParagraphStyle('code',fontName='Chinese',fontSize=8.1,leading=13,wordWrap='CJK',textColor=NAVY,backColor=colors.HexColor('#EEF4F5'),borderPadding=8,spaceAfter=10),
 'cover':ParagraphStyle('cover',fontName='Chinese',fontSize=29,leading=42,textColor=NAVY,wordWrap='CJK'),
 'sub':ParagraphStyle('sub',fontName='Chinese',fontSize=14,leading=23,textColor=TEAL,spaceAfter=18,wordWrap='CJK'),
 'small':ParagraphStyle('small',fontName='Chinese',fontSize=9,leading=14,textColor=GRAY,wordWrap='CJK'),
}

def clean(s):
    s=re.sub(r'\[([^\]]+)\]\([^)]*\)',r'\1',s)
    return s.replace('**','').replace('`','')

def para(s,style='body'):
    return Paragraph(html.escape(clean(s)).replace('\n','<br/>'),styles[style])

def weights(headers):
    n=len(headers)
    if n==6 and headers[0]=='实体脚': return [0.055,0.12,0.22,0.17,0.215,0.22]
    if n==6 and headers[0]=='槽位': return [0.055,0.15,0.23,0.10,0.205,0.26]
    if n==6 and headers[0]=='板' and headers[2]=='数': return [0.055,0.18,0.055,0.315,0.13,0.265]
    if n==6 and headers[0]=='板' and headers[2]=='型号': return [0.055,0.08,0.245,0.285,0.25,0.085]
    if n==7: return [0.15,0.1,0.075,0.085,0.085,0.105,0.4]
    if n==6 and headers[0]=='编号': return [0.065,0.245,0.21,0.20,0.10,0.18]
    if n==5 and headers[0]=='焊盘': return [0.07,0.12,0.25,0.49,0.07]
    if n==6 and headers[1]=='信号': return [0.05,0.16,0.23,0.25,0.23,0.08]
    if n==3: return [0.21,0.30,0.49]
    if n==4: return [0.16,0.24,0.30,0.30]
    return [1/n]*n

class Handbook(BaseDocTemplate):
    def __init__(self,filename,**kw):
        super().__init__(filename,pagesize=A4,leftMargin=17*mm,rightMargin=17*mm,topMargin=19*mm,bottomMargin=18*mm,**kw)
        self.addPageTemplates(PageTemplate(id='A4',frames=[Frame(self.leftMargin,self.bottomMargin,self.width,self.height,id='normal',leftPadding=0,rightPadding=0,topPadding=0,bottomPadding=0)],onPage=self.header_footer))
    def header_footer(self,canvas,doc):
        canvas.saveState()
        canvas.setFont('Chinese',7.5); canvas.setFillColor(GRAY)
        if doc.page>1:
            canvas.drawString(17*mm,A4[1]-11*mm,'四轮四转底盘 · VCU / BCM 硬件实施手册')
            canvas.drawRightString(A4[0]-17*mm,A4[1]-11*mm,'HW-R1 | 2026-09-19')
            canvas.setStrokeColor(colors.HexColor('#C9D8DD')); canvas.line(17*mm,A4[1]-14*mm,A4[0]-17*mm,A4[1]-14*mm)
        canvas.setFont('Chinese',7)
        canvas.drawString(17*mm,10*mm,'设计输入版 · 未经硬件验证 · 不得直接投板')
        canvas.drawRightString(A4[0]-17*mm,10*mm,str(doc.page))
        canvas.restoreState()
    def afterFlowable(self,flowable):
        if hasattr(flowable,'section_level'):
            level=flowable.section_level
            self.canv.bookmarkPage(flowable.bookmark)
            self.canv.addOutlineEntry(flowable.getPlainText(),flowable.bookmark,level=level,closed=(level>0))
            if level==0: self.notify('TOCEntry',(0,flowable.getPlainText(),self.page,flowable.bookmark))

pdf_path=OUT/'四轮四转底盘_硬件实施与检验手册_HW-R1.pdf'
story=[Spacer(1,28*mm),para('四轮四转底盘\n硬件实施与检验手册','cover'),Spacer(1,10*mm),
       para('VCU + BCM\n原理图输入 · 线束装配 · PCB 制造 · 样板检验','sub'),
       para('HW-R1 / 2026-09-19','sub'),Spacer(1,8*mm),
       para('包含 228 个 MCU 实体脚、310 条核心芯片焊盘/EP 连接、35P 逻辑接口、12 路传感器映射、参考电路、计算与 26 项检查程序。'),
       Spacer(1,7*mm),para('发布状态：设计输入版','h2'),
       para('已核验资料、工程建议和待确认项分别标识。连接器厂家腔号、真实外设型号与安全电路等 H 项关闭前，不可直接画成生产定稿或交厂投板。'),
       para('本手册不包含已验证 CAD/Gerber、可采购完整 BOM 或实际测试合格记录。纸面校验不能代替样板与整车验证。'),
       Spacer(1,14*mm),para('使用人：________    项目版本：________\n设计复核：________    领取日期：________','small'),PageBreak(),para('目录','h1')]
toc=TableOfContents()
toc.levelStyles=[ParagraphStyle('toc',fontName='Chinese',fontSize=10,leading=19,spaceAfter=4,wordWrap='CJK')]
story.extend([toc,Spacer(1,10*mm),para('建议 A4 原尺寸打印，长边装订。页眉版本必须与线束、原理图和工单一致。','small')])
counter=0
for block_index,(kind,content,source) in enumerate(blocks):
    if kind.startswith('h'):
        if kind=='h1': story.append(PageBreak())
        p=para(content,kind)
        # Long tables must be allowed to begin below a heading on the same page.
        # Keeping a whole multi-page table with its heading wastes most of a page.
        if kind in ['h2','h3'] and block_index+1<len(blocks) and blocks[block_index+1][0]=='table' and len(blocks[block_index+1][1])>8:
            story.append(CondPageBreak(95))
            p.keepWithNext=False
        if kind in ['h1','h2']:
            p.section_level=int(kind[1])-1
            p.bookmark='sec'+str(counter); counter+=1
        story.append(p)
    elif kind=='table':
        cellrows=[[para(cell,'thead' if j==0 else 'cell') for cell in row] for j,row in enumerate(content)]
        tbl=LongTable(cellrows,colWidths=[x*(A4[0]-34*mm) for x in weights(content[0])],repeatRows=1,hAlign='LEFT')
        tbl.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),TEAL),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F0F5F6')]),('VALIGN',(0,0),(-1,-1),'TOP'),('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#CAD7DC')),('TOPPADDING',(0,0),(-1,-1),5),('BOTTOMPADDING',(0,0),(-1,-1),5),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5)]))
        story.extend([tbl,Spacer(1,7)])
    elif kind=='code': story.append(para(content,'code'))
    else: story.append(para(content))
Handbook(str(pdf_path),title='四轮四转底盘硬件实施与检验手册 HW-R1',author='项目硬件设计资料',subject='设计输入；未通过硬件/生产放行').multiBuild(story)

# Editable companion, with repeated table headers and explicit row preservation.
doc=Document()
section=doc.sections[0]
section.page_width=Mm(210); section.page_height=Mm(297)
section.top_margin=Mm(20); section.bottom_margin=Mm(18)
section.left_margin=Mm(17); section.right_margin=Mm(17)
for style in doc.styles:
    if style.type==1:
        style.font.name='Arial Unicode MS'
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
normal=doc.styles['Normal']; normal.font.size=Pt(10); normal.paragraph_format.space_after=Pt(6)
normal.paragraph_format.line_spacing=1.18
section.header.paragraphs[0].text='四轮四转底盘 · 硬件实施与检验手册 | HW-R1'
footer=section.footer.paragraphs[0]; footer.text='设计输入版 · 未经硬件验证 · 不得直接投板    页 '
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.add_heading('四轮四转底盘\n硬件实施与检验手册',0)
doc.add_paragraph('VCU + BCM | HW-R1 | 2026-09-19')
doc.add_paragraph('可编辑源文件。正式打印优先使用同版本 PDF；编辑本文件后必须回写 Markdown/结构化数据并重新生成，避免版本分叉。')
doc.add_paragraph('发布状态：设计输入版。所有 H 项必须关闭，实际 CAD 与样板验证完成后才能生产放行。')
doc.add_paragraph('使用人：________  项目版本：________\n独立复核：________  日期：________')
doc.add_page_break()
doc.add_heading('章节索引',1)
for kind,content,source in blocks:
    if kind=='h1': doc.add_paragraph(content)
for kind,content,source in blocks:
    if kind.startswith('h'):
        if kind=='h1': doc.add_page_break()
        doc.add_heading(clean(content),int(kind[1]))
    elif kind=='table':
        t=doc.add_table(rows=1,cols=len(content[0])); t.style='Table Grid'; t.autofit=False
        for idx,w in enumerate(weights(content[0])): t.columns[idx].width=Mm(176*w)
        for idx,row in enumerate(content):
            cells=t.rows[0].cells if idx==0 else t.add_row().cells
            for ci,(cell,txt) in enumerate(zip(cells,row)):
                cell.width=Mm(176*weights(content[0])[ci]); cell.text=clean(txt)
                for p in cell.paragraphs:
                    p.paragraph_format.space_after=Pt(3); p.paragraph_format.line_spacing=1.1
                    for run in p.runs: run.font.size=Pt(8)
                if idx==0:
                    shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'DDEDEF'); cell._tc.get_or_add_tcPr().append(shd)
            trPr=t.rows[idx]._tr.get_or_add_trPr(); trPr.append(OxmlElement('w:cantSplit'))
            if idx==0: trPr.append(OxmlElement('w:tblHeader'))
        doc.add_paragraph('')
    elif kind=='code':
        p=doc.add_paragraph(clean(content))
        for r in p.runs: r.font.size=Pt(8)
    else: doc.add_paragraph(clean(content))
docx_path=OUT/'四轮四转底盘_硬件实施与检验手册_HW-R1.docx'
doc.save(docx_path)

reader=PdfReader(str(pdf_path))
page_texts=[p.extract_text() or '' for p in reader.pages]
text='\n'.join(page_texts)
required=['MAX14830','MC33926','TPSM53603','TCAN1044AV','128','100','35P','T26','H16','SHA256','HOLD_VERIFY_REF']
for word in required: assert word in text,('PDF missing',word)
assert all(len(t.strip())>15 for t in page_texts),'Unexpected blank PDF page'
assert '禁止直接投板' not in text or len(reader.pages)>10
report={'status':'EXPORT_CONTENT_CHECK_PASSED_NOT_HARDWARE_VALIDATION','pages':len(reader.pages),'source_files':[p.name for p in FILES],'tables':sum(k=='table' for k,_,_ in blocks),'required_terms_checked':required,'artifacts':[],'visual_review':'Required: rendered-page review is separate from text/content checks.'}
for path in [pdf_path,docx_path]:
    report['artifacts'].append({'file':path.name,'bytes':path.stat().st_size,'sha256':hashlib.sha256(path.read_bytes()).hexdigest()})
(OUT/'导出检查.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(report,ensure_ascii=False,indent=2))
